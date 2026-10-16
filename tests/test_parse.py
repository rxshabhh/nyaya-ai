# Parser checks for txt, digital pdf, docx and a scanned (textless) pdf.
# Run: python tests/test_parse.py
import sys
import tempfile
from pathlib import Path

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT))

from docx import Document
from docx.enum.text import WD_BREAK

from nyaya.parse import ParseError, load_para


def make_pdf(pages):
    # pages: list of pages, each a list of (y, text); an empty list is a blank (scanned-like) page
    objs = ["<< /Type /Catalog /Pages 2 0 R >>", None, "<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>"]
    kids = []
    for lines in pages:
        stream = "BT /F1 12 Tf " + " ".join(f"1 0 0 1 72 {y} Tm ({t}) Tj" for y, t in lines) + " ET"
        n = len(objs) + 1
        kids.append(f"{n} 0 R")
        objs.append(f"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] /Contents {n + 1} 0 R "
                    "/Resources << /Font << /F1 3 0 R >> >> >>")
        objs.append(f"<< /Length {len(stream)} >>\nstream\n{stream}\nendstream")
    objs[1] = f"<< /Type /Pages /Kids [{' '.join(kids)}] /Count {len(kids)} >>"
    out, offsets = "%PDF-1.4\n", []
    for i, o in enumerate(objs, 1):
        offsets.append(len(out))
        out += f"{i} 0 obj\n{o}\nendobj\n"
    xref = len(out)
    out += f"xref\n0 {len(objs) + 1}\n0000000000 65535 f \n"
    out += "".join(f"{o:010d} 00000 n \n" for o in offsets)
    out += f"trailer\n<< /Size {len(objs) + 1} /Root 1 0 R >>\nstartxref\n{xref}\n%%EOF\n"
    return out.encode("latin-1")


with tempfile.TemporaryDirectory() as tmp:
    tmp = Path(tmp)
    (tmp / "a.txt").write_text("[page 1]\n\nFirst para.\n\nSecond para.\n\n[page 2]\n\nThird.\n", encoding="utf-8")
    (tmp / "b.pdf").write_bytes(make_pdf([
        [(700, "On 12.03.2021 at about 9 pm I was at my house"), (686, "when the accused Suresh came there."),
         (650, "He was carrying a knife.")],
        [(700, "Second page text.")]]))
    d = Document()
    d.add_paragraph("Docx one.")
    d.add_paragraph("Docx two.").add_run().add_break(WD_BREAK.PAGE)
    d.add_paragraph("Docx three.")
    d.save(tmp / "c.docx")
    (tmp / "notes.md").write_text("ignored", encoding="utf-8")

    paras = load_para(tmp)
    assert paras[("a.txt", 1, 2)] == "Second para." and paras[("a.txt", 2, 1)] == "Third."
    assert paras[("b.pdf", 1, 1)] == "On 12.03.2021 at about 9 pm I was at my house when the accused Suresh came there."
    assert paras[("b.pdf", 1, 2)] == "He was carrying a knife." and paras[("b.pdf", 2, 1)] == "Second page text."
    assert paras[("c.docx", 1, 2)] == "Docx two." and paras[("c.docx", 2, 1)] == "Docx three."
    assert not any(k[0] == "notes.md" for k in paras)

    (tmp / "b.pdf").write_bytes(make_pdf([[(700, "Has text.")], []]))
    try:
        load_para(tmp)
        raise SystemExit("scanned pdf was accepted")
    except ParseError as e:
        assert "page(s) 2" in str(e), e

print("ok")
